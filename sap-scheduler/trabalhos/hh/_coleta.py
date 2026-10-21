# -*- coding: utf-8 -*-
"""Coleta SAP de HH (IW37 planejado / IW47 realizado) — navegação + export ALV nativo.

Promovido do protótipo `exploratorio/iw_fastexport.py`. READ-ONLY no SAP: navega,
adiciona todas as colunas (move-all no layout — só exibição, não salva variante),
dispara o export &XXL → .xlsx (um disparo), lê o xlsx e mapeia as colunas-alvo por
título de exibição. Reusa `zpp_gridcap.export_grid` (VIA C provada).

A unidade (`ARBEH`/`ISMNE`) chega em MIN (campanha validacao-coleta-e-consumo.md §4);
a agregação fica em `_agg.py`, o parsing em `_parse.py`.
"""
from __future__ import annotations

import os
import tempfile
import time
from types import SimpleNamespace

from ..kpis_manut import zpp_gridcap as z

# Campos-alvo por transação (do protótipo iw_fastexport.py, validados na campanha)
TARGET_IW47 = ["AUFNR", "VORNR", "WERKI", "AUART", "STTXT", "PERNR", "NAME", "TPLNR",
               "ISMNW", "ISMNE", "ARBEI", "ARBEH", "PARBPL", "IARBPL",
               "ISDD", "ISDZ", "IEDD", "IEDZ", "ERSDA", "BUDAT",
               "STOKZ", "STZHL", "RUECK", "RMZHL"]
TARGET_IW37 = ["AUFNR", "VORNR", "AUART", "LTXA1", "VSTTXT", "PRIOK", "ARBPL", "IWERK", "WERKS",
               "TPLNR", "EQUNR", "EQKTX", "ARBEI", "ARBEH",
               "FSAVD", "FSAVZ", "FSEDD", "FSEDZ", "SSAVD", "SSAVZ", "SSEDD", "SSEDZ",
               "ISDD", "IEDD", "PERNR"]


def _drain_popups(s) -> None:
    """Fecha popups perdidos (F12 em wnd[2]/wnd[1]) antes de navegar — read-only."""
    for w in ("wnd[2]", "wnd[1]"):
        try:
            s.findById(w).sendVKey(12)
            time.sleep(0.3)
        except Exception:
            pass


def _clip_set(vals) -> None:
    """Copia valores (um por linha) pro clipboard — p/ o Upload da multisseleção SAP.

    Win-only, importado sob demanda (não quebra o import no Linux/testes, que usam o mock).
    """
    import win32clipboard  # noqa: PLC0415 — Windows-only, sob demanda
    win32clipboard.OpenClipboard()
    win32clipboard.EmptyClipboard()
    win32clipboard.SetClipboardText("\r\n".join(vals))
    win32clipboard.CloseClipboard()


def navigate(s, tx: str, low: str, high: str) -> None:
    """Entra em /n<tx>, aplica filtro de planta BR02 + janela de datas, F8. Read-only."""
    s.findById("wnd[0]/tbar[0]/okcd").Text = f"/n{tx}"
    s.findById("wnd[0]").sendVKey(0)
    time.sleep(0.8)
    if tx == "iw47":
        for k, v in (("ctxtVARIANT", "/TRAB REAL"), ("ctxtWERKS_C-LOW", "BR02"),
                     ("ctxtERSDA_C-LOW", low), ("ctxtERSDA_C-HIGH", high)):
            try:
                s.findById("wnd[0]/usr/" + k).Text = v
            except Exception:
                pass
    else:  # iw37
        for chk in ("chkDY_AKT", "chkDY_HIS"):
            try:
                s.findById("wnd[0]/usr/" + chk).selected = True
            except Exception:
                pass
        for k, v in (("ctxtWERKS-LOW", "BR02"), ("ctxtFSAVD-LOW", low), ("ctxtFSAVD-HIGH", high)):
            try:
                s.findById("wnd[0]/usr/" + k).Text = v
            except Exception:
                pass
    s.findById("wnd[0]").sendVKey(8)
    time.sleep(1.8)


def _open_layout_and_move_all(s) -> bool:
    """Abre o diálogo de layout 'Atuais...' e move TODAS as colunas para a exibição.

    Garante que as colunas-alvo (ARBEI/ARBEH/STTXT/PERNR…) estejam no grid antes do
    export. Só exibição — não salva variante (read-only).
    """
    found = []

    def walk(node):
        try:
            kids = node.Children
            n = kids.Count
        except Exception:
            return
        for i in range(n):
            try:
                ch = kids(i)
            except Exception:
                continue
            if "atuais" in (getattr(ch, "Text", "") or "").lower():
                found.append(getattr(ch, "Id", ""))
            walk(ch)

    walk(s.findById("wnd[0]/mbar"))
    if not found:
        return False
    s.findById(found[0]).select()
    time.sleep(1.0)
    base = "wnd[1]/usr/tabsG_TS_ALV/tabpALV_M_R1/ssubSUB_DYN0510:SAPLSKBH:0620/"
    setgrid = s.findById(base + "cntlCONTAINER1_LAYO/shellcont/shell")
    n = setgrid.RowCount
    setgrid.selectedRows = f"0-{n-1}"
    s.findById(base + "btnAPP_WL_SING").press()
    time.sleep(0.8)
    s.findById("wnd[1]/tbar[0]/btn[0]").press()
    time.sleep(1.5)
    return True


def _export_and_read(s, tx: str) -> list[dict]:
    """Exporta o grid (&XXL → .xlsx) e lê, mapeando colunas-alvo por título de exibição.

    Replica o alinhamento sequencial do protótipo: o &XXL exporta na ordem do grid mas
    descarta colunas de controle; percorre o header do xlsx avançando no ColumnOrder até
    casar o título. Retorna list[dict] com chave = nome técnico SAP.
    """
    grid = z.find_grid(s.findById("wnd[0]/usr"))
    if grid is None:
        raise RuntimeError("hh: grid não encontrado na tela (navegação falhou)")
    cols_tech = [str(c) for c in z.coll(grid.ColumnOrder) if c is not None]
    disp_tech = [(c, z.col_title(grid, c)) for c in cols_tech]

    save_dir = tempfile.gettempdir()
    fname = f"hh_{tx}_viaC.xlsx"
    xlsx = os.path.join(save_dir, fname)
    args = SimpleNamespace(tx=None, set=[], no_f8=False, ctx_code="&XXL",
                           kill_excel=True, save_dir=save_dir, fname=fname)
    z.export_grid(s, args)
    if not os.path.exists(xlsx):
        raise RuntimeError(f"hh: xlsx não gerado: {xlsx}")

    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_disp = list(next(rows_iter))
    # alinhamento sequencial título-de-exibição → nome técnico
    tech_by_pos = [None] * len(header_disp)
    ti = 0
    for xi, h in enumerate(header_disp):
        h = h or ""
        start = ti
        while ti < len(disp_tech) and (disp_tech[ti][1] or "") != h:
            ti += 1
        if ti < len(disp_tech):
            tech_by_pos[xi] = disp_tech[ti][0]
            ti += 1
        else:
            ti = start
    target = TARGET_IW47 if tx == "iw47" else TARGET_IW37
    idx = {t: tech_by_pos.index(t) for t in target if t in tech_by_pos}

    out = []
    for row in rows_iter:
        out.append({t: row[idx[t]] for t in idx})
    try:
        wb.close()
        os.remove(xlsx)
    except Exception:
        pass
    return out


def coletar(session, tx: str, low: str, high: str) -> list[dict]:
    """Coleta completa de uma transação: drena popups, navega, move colunas, exporta+lê.

    Usado pelo `hh_realizado` (IW47, janela `ERSDA` recente). O `hh_planejado` NÃO usa mais
    este caminho — desde o B14 ele coleta via `coletar_planejado` (IW37 por multisseleção de
    AUFNR, sem janela FSAVD), pois a janela ampla estoura dump ABAP neste sistema.
    """
    _drain_popups(session)
    navigate(session, tx, low, high)
    if not _open_layout_and_move_all(session):
        raise RuntimeError("hh: não consegui abrir o layout 'Atuais...' para mover colunas")
    return _export_and_read(session, tx)


def navigate_planejado(s, aufnr_list) -> None:
    """IW37 restrito às AUFNR abertas (multisseleção via clipboard), SEM janela FSAVD.

    B14 / SP-01.5 revisado: o IW37 sem restringir o CONJUNTO de ordens estoura dump ABAP
    "memória escassa" neste sistema (o gargalo é o volume de operações "Atual", não o span
    de data). A multisseleção por AUFNR é o que limita o volume. `DY_AKT`+`DY_PMONL` on,
    `DY_HIS` off, planta BR02. Read-only. Mesmo padrão de `backlog._coleta._iw37_centro_op0010`.
    """
    s.findById("wnd[0]/tbar[0]/okcd").Text = "/niw37"
    s.findById("wnd[0]").sendVKey(0)
    time.sleep(1.0)
    # limpa campos que restringiriam/contradiriam a multisseleção (datas e ranges de ordem)
    for idp in ("ctxtAUFNR-LOW", "ctxtAUFNR-HIGH", "ctxtFSAVD-LOW", "ctxtFSAVD-HIGH",
                "ctxtDATUV", "ctxtDATUB"):
        try:
            s.findById("wnd[0]/usr/" + idp).Text = ""
        except Exception:
            pass
    try:
        s.findById("wnd[0]/usr/ctxtWERKS-LOW").Text = "BR02"
    except Exception:
        pass
    # multisseleção de AUFNR: upload do clipboard (nunca faixa min-máx — elimina ordens do meio)
    _clip_set(list(aufnr_list))
    s.findById("wnd[0]/usr/ctxtAUFNR-LOW").setFocus()
    s.findById("wnd[0]/usr/btn%_AUFNR_%_APP_%-VALU_PUSH").press()
    time.sleep(0.8)
    s.findById("wnd[1]/tbar[0]/btn[24]").press()   # aba "Upload from clipboard"
    time.sleep(0.8)
    s.findById("wnd[1]/tbar[0]/btn[8]").press()     # executar seleção
    time.sleep(0.8)
    for cb, v in (("DY_AKT", True), ("DY_PMONL", True), ("DY_HIS", False)):
        try:
            s.findById("wnd[0]/usr/chk" + cb).Selected = v
        except Exception:
            pass
    s.findById("wnd[0]").sendVKey(8)
    time.sleep(2.5)
    try:
        if s.findById("wnd[1]", False) is not None:
            s.findById("wnd[1]").sendVKey(0)   # confirma popup de "seleção grande", se surgir
            time.sleep(0.5)
    except Exception:
        pass


def coletar_planejado(session, aufnr_list) -> list[dict]:
    """Coleta IW37 (planejado/`ARBEI`) restrita às AUFNR abertas do IW39. Read-only.

    Substitui `coletar(session, 'iw37', low, high)` no `run_planejado` (B14): sem janela
    FSAVD, com multisseleção por ordem. Reusa o move-all + export ALV nativo já validados.
    """
    _drain_popups(session)
    navigate_planejado(session, aufnr_list)
    if not _open_layout_and_move_all(session):
        raise RuntimeError("hh: não consegui abrir o layout 'Atuais...' para mover colunas")
    return _export_and_read(session, "iw37")


def navigate_realizado_por_aufnr(s, aufnr_list) -> None:
    """IW47 restrito às AUFNR-alvo (multisseleção via clipboard), SEM janela de data (BR-20).

    Usado pela re-consulta dirigida das ordens que SAÍRAM do backlog (concluídas): pega o
    realizado final + status (`STTXT`) dessas ordens, mesmo que a confirmação seja antiga (por
    isso sem janela `ERSDA` — a multisseleção por ordem é o que limita o volume). Variante
    `/TRAB REAL`, planta BR02. Read-only.

    ATENÇÃO ao campo da ordem: no **IW47** o campo/botão de ordem é **`AUFNR_O`** (não `AUFNR`
    como no IW37) — validado no cliente 2026-07-14 (GM3/040): IDs reais `ctxtAUFNR_O-LOW` e
    `btn%_AUFNR_O_%_APP_%-VALU_PUSH`. Ver `sap-gate/b26-validacao/b26_relatorio.json`."""
    s.findById("wnd[0]/tbar[0]/okcd").Text = "/niw47"
    s.findById("wnd[0]").sendVKey(0)
    time.sleep(1.0)
    for k, v in (("ctxtVARIANT", "/TRAB REAL"), ("ctxtWERKS_C-LOW", "BR02")):
        try:
            s.findById("wnd[0]/usr/" + k).Text = v
        except Exception:
            pass
    # limpa a janela de data (queremos TODAS as confirmações das ordens-alvo, não só recentes)
    for idp in ("ctxtERSDA_C-LOW", "ctxtERSDA_C-HIGH", "ctxtAUFNR_O-LOW", "ctxtAUFNR_O-HIGH"):
        try:
            s.findById("wnd[0]/usr/" + idp).Text = ""
        except Exception:
            pass
    # multisseleção de AUFNR: upload do clipboard (campo `AUFNR_O` no IW47 — validado no cliente)
    _clip_set(list(aufnr_list))
    s.findById("wnd[0]/usr/ctxtAUFNR_O-LOW").setFocus()
    s.findById("wnd[0]/usr/btn%_AUFNR_O_%_APP_%-VALU_PUSH").press()
    time.sleep(0.8)
    s.findById("wnd[1]/tbar[0]/btn[24]").press()   # aba "Upload from clipboard"
    time.sleep(0.8)
    s.findById("wnd[1]/tbar[0]/btn[8]").press()     # executar seleção
    time.sleep(0.8)
    s.findById("wnd[0]").sendVKey(8)
    time.sleep(2.0)
    try:
        if s.findById("wnd[1]", False) is not None:
            s.findById("wnd[1]").sendVKey(0)
            time.sleep(0.5)
    except Exception:
        pass


def coletar_realizado_por_aufnr(session, aufnr_list) -> list[dict]:
    """Coleta IW47 (realizado final + `STTXT`) restrita às AUFNR-alvo. Read-only (BR-20/SP-09.6b).

    Re-consulta dirigida das ordens que saíram do backlog: sem janela de data, com multisseleção
    por ordem. Reusa o move-all + export ALV nativo já validados.

    ATENÇÃO ao resultado VAZIO (diag 2026-07-14): ordem cancelada/TECO sem apontamento não tem
    confirmação no IW47 → o F8 por AUFNR_O NÃO renderiza ALV. Nesse caso o job NÃO deve estourar:
    devolve `[]` (sem realizado) e o `run_entregas` resolve o status dessas ordens por outro caminho
    (IW39-por-AUFNR). Só levanta se HÁ grid mas o menu de layout não abre (aí sim é anomalia real).

    O move-all às vezes falha ("não achei o menu Atuais") por popup residual ou porque o grid ainda
    não terminou de pintar quando o walk do menu roda. Por isso: drena popup + espera crescente +
    retenta até 3×; a cada tentativa checa se o grid já apareceu."""
    _drain_popups(session)
    navigate_realizado_por_aufnr(session, aufnr_list)
    grid_visto = False
    for tentativa in range(3):
        _drain_popups(session)                 # fecha popup que bloqueia o menu de layout
        time.sleep(1.0 + tentativa)            # dá tempo do grid pintar (1s, 2s, 3s)
        if z.find_grid(session.findById("wnd[0]/usr")) is None:
            continue                           # ainda sem ALV — pode ser 0 confirmação OU pintando
        grid_visto = True
        if _open_layout_and_move_all(session):
            return _export_and_read(session, "iw47")
    if not grid_visto:
        # nenhuma confirmação p/ as ordens-alvo (canceladas/TECO/sem apontamento) → sem realizado.
        return []
    raise RuntimeError("hh: layout 'Atuais...' não abriu apesar do grid presente "
                       "(após 3 tentativas com drenagem de popup)")


def coletar_status_por_aufnr(session, aufnr_list) -> dict:
    """Status de ordem (STTXT/USTXT/AUART/TPLNR) por AUFNR via IW39, TODOS os status (BR-20/SP-09.6d).

    Usado só p/ as ordens que saíram do backlog mas voltaram VAZIAS do IW47 (canceladas/TECO sem
    apontamento): o realizado não existe, mas o status verdadeiro (CANC/ENTE...) está no IW39. IW39
    por multisseleção de AUFNR (mesmos IDs do backlog), com DY_OFN+DY_IAR+DY_MAB+DY_HIS ligados p/
    incluir concluídas/canceladas/histórico. Reusa os helpers do backlog (grid/export/expor colunas
    — o /BACKLOG não traz USTXT, então força move-all). Read-only.

    FAIL-SOFT por design: QUALQUER erro/tela inesperada → `{}` (o chamador trata como 'não resolvido'
    e a ordem reentra no próximo ciclo). Nunca derruba o job `hh_entregas`. Devolve
    {aufnr: {status_sistema, status_usuario, auart, tplnr}}.
    """
    from ..backlog import _coleta as bl  # noqa: PLC0415 — reuso de find_grid/coll/export/expor
    import os as _os
    import tempfile as _tf
    s = session
    try:
        s.findById("wnd[0]/tbar[0]/okcd").Text = "/niw39"
        s.findById("wnd[0]").sendVKey(0)
        time.sleep(1.0)
        for idp in ("ctxtAUFNR-LOW", "ctxtAUFNR-HIGH", "ctxtGEWRK-LOW", "ctxtARBPL-LOW",
                    "ctxtAUART-LOW", "ctxtAUART-HIGH", "ctxtSTAE1-LOW", "ctxtSTAE1-HIGH", "ctxtSTAI1-LOW"):
            try:
                s.findById("wnd[0]/usr/" + idp).Text = ""
            except Exception:
                pass
        try:
            s.findById("wnd[0]/usr/ctxtVARIANT").Text = bl.VARIANTE
        except Exception:
            pass
        s.findById("wnd[0]/usr/ctxtIWERK-LOW").Text = "BR02"
        try:
            s.findById("wnd[0]/usr/ctxtDATUV").Text = "01.01.2010"
            s.findById("wnd[0]/usr/ctxtDATUB").Text = "31.12.2035"
        except Exception:
            pass
        # multisseleção de AUFNR (mesmo botão do IW37): upload do clipboard
        _clip_set([str(a) for a in aufnr_list])
        s.findById("wnd[0]/usr/ctxtAUFNR-LOW").setFocus()
        s.findById("wnd[0]/usr/btn%_AUFNR_%_APP_%-VALU_PUSH").press()
        time.sleep(0.6)
        s.findById("wnd[1]/tbar[0]/btn[24]").press()
        time.sleep(0.5)
        s.findById("wnd[1]/tbar[0]/btn[8]").press()
        time.sleep(0.5)
        # TODOS os status (pega cancelada/TECO/concluída/histórico), diferente do backlog (só abertas)
        for cb, v in {"DY_OFN": True, "DY_IAR": True, "DY_MAB": True, "DY_HIS": True}.items():
            try:
                s.findById("wnd[0]/usr/chk" + cb).Selected = v
            except Exception:
                pass
        s.findById("wnd[0]").sendVKey(8)
        time.sleep(2.0)
        try:
            if s.findById("wnd[1]", False) is not None:
                s.findById("wnd[1]").sendVKey(0)
                time.sleep(0.5)
        except Exception:
            pass
        g = bl._find_grid(s)
        if g is None:
            return {}
        cols = [c for c in bl._coll(g.ColumnOrder) if c is not None]
        if "USTXT" not in cols or "STTXT" not in cols:
            g2 = bl._expor_todas_colunas(s)   # /BACKLOG não traz USTXT → força move-all
            if g2 is not None:
                g = g2
        target = _os.path.join(_tf.gettempdir(), "hh_status_aufnr.xlsx")
        # KTEXT = texto breve da ordem (título). O IW47 não traz KTEXT nas confirmações (só o IW39),
        # então é aqui que o título da ordem que saiu do backlog é recuperado (senão o card perde o título).
        recs = bl._export_and_read(s, g, ["AUFNR", "STTXT", "USTXT", "AUART", "TPLNR", "KTEXT"], target)
        out: dict = {}
        for r in recs:
            a = (r.get("AUFNR") or "").strip()
            if a and a not in out:
                out[a] = {"status_sistema": (r.get("STTXT") or "").strip() or None,
                          "status_usuario": (r.get("USTXT") or "").strip() or None,
                          "auart": (r.get("AUART") or "").strip() or None,
                          "tplnr": (r.get("TPLNR") or "").strip() or None,
                          "descricao": (r.get("KTEXT") or "").strip() or None}
        return out
    except Exception:
        return {}


# ---- Mock (teste sem SAP — ctx.cfg.mock_sap) -----------------------------------------

def mock_planejado() -> list[dict]:
    """Linhas IW37 sintéticas (com linha de total) para teste do job sem SAP."""
    return [
        {"AUFNR": "5018863", "VORNR": "0010", "ARBEI": "120", "ARBEH": "MIN",
         "ARBPL": "Z210", "FSAVD": "01.06.2026"},
        {"AUFNR": "5018863", "VORNR": "0010", "ARBEI": "120", "ARBEH": "MIN",
         "ARBPL": "Z210", "FSAVD": "01.06.2026"},   # repete → dedup
        {"AUFNR": "5018901", "VORNR": "0010", "ARBEI": "480", "ARBEH": "MIN",
         "ARBPL": "Z110", "FSAVD": "02.06.2026"},
        {"AUFNR": "", "VORNR": "", "ARBEI": "600", "ARBEH": "MIN"},   # linha de total
    ]


def mock_realizado() -> list[dict]:
    """Linhas IW47 sintéticas (com linha de total) para teste do job sem SAP."""
    return [
        {"AUFNR": "5018863", "VORNR": "0010", "ISMNW": "60", "ISMNE": "MIN",
         "PERNR": "CDAMASIO", "NAME": "Carlos Damasio", "STTXT": "LIB ENCE"},
        {"AUFNR": "5018863", "VORNR": "0010", "ISMNW": "50", "ISMNE": "MIN",
         "PERNR": "CDAMASIO", "NAME": "Carlos Damasio", "STTXT": "LIB ENCE"},
        {"AUFNR": "5018901", "VORNR": "0010", "ISMNW": "500", "ISMNE": "MIN",
         "PERNR": "JSILVA", "NAME": "Joao Silva", "STTXT": "LIB"},
        {"AUFNR": "", "VORNR": "", "ISMNW": "610", "ISMNE": "MIN"},   # linha de total
    ]


def mock_status_por_aufnr(aufnr_list) -> dict:
    """Mock do IW39-por-AUFNR: status do usuário `CERR` (Encerrada) p/ o e2e exercitar a resolução
    terminal por status do usuário (CERR→encerrada). Cabeçalho parcial (AUART/TPLNR) como no real."""
    return {str(a): {"status_usuario": "CERR", "status_sistema": "ENTE", "auart": "YPM1",
                     "tplnr": "BR02-0801", "descricao": "Troca rolamento"} for a in aufnr_list}


def mock_realizado_por_aufnr(aufnr_list) -> list[dict]:
    """Linhas IW47 sintéticas para a re-consulta dirigida (BR-20) sem SAP: devolve o realizado
    final + status 'concluído' das AUFNR-alvo. Uma linha por ordem (o `_agg` soma/deriva)."""
    rows = []
    for i, a in enumerate(aufnr_list):
        rows.append({"AUFNR": str(a), "VORNR": "0010", "ISMNW": str(100 + i), "ISMNE": "MIN",
                     "ARBEI": str(120 + i), "ARBEH": "MIN",   # IW47 traz o planejado (TARGET_IW47)
                     "AUART": "YPM1", "TPLNR": "BR02-0801", "PARBPL": "Z12",   # cabeçalho (B29)
                     "PERNR": "CDAMASIO", "NAME": "Carlos Damasio", "STTXT": "LIB ENCE ENTG"})
    rows.append({"AUFNR": "", "VORNR": "", "ISMNW": str(sum(100 + i for i in range(len(aufnr_list)))),
                 "ISMNE": "MIN"})   # linha de total
    return rows
