# -*- coding: utf-8 -*-
"""Testes da planilha do planejamento semanal (hh-management B55).

Exercita a camada PURA `montar_planilha` (abas, realizado rateado sem dupla contagem, agregação
por ordem, fila a planejar, replanejamentos) e o `gerar_xlsx` (bytes de .xlsx válido, relido com
openpyxl) — sem Mongo; o quadro é um mock no formato de `hh_data.montar_quadro`.
"""
import io
from datetime import datetime

from openpyxl import load_workbook

from src.utils import hh_data
from src.utils import plano_semanal_xlsx as px

SEMANA = "2026-W38"
EXPORTADO_EM = datetime(2026, 9, 18, 10, 30)


def _ordem(num, plan, real=0, parcela=None, atrib_id=None, **extra):
    """Card de ordem no formato do quadro (parcela None = ordem inteira)."""
    o = {"ordem": num, "descricao": f"ordem {num}", "disciplina": "Mecanica",
         "planejado_min": plan, "realizado_min": real,
         "aderencia": (real / plan) if plan else None, "parcela_min": parcela,
         "_atrib_id": atrib_id or f"a-{num}-{parcela}", "_status_usuario_label": "Liberada",
         "linha_nome": "Linha 1", "status_sap": "LIB", "executores": ["JOAO"], "operacoes": 2}
    o.update(extra)
    return o


def _tec(fid, nome, dias, turno="A"):
    base = {d: [] for d in range(7)}
    base.update(dias)
    return {"funcionario_id": fid, "nome": nome, "turno": turno, "dias": base}


def _quadro(mec=None, ele=None, a_planejar=None, terceiros=None, pendentes=None):
    cap = {"A": {d: {"disponivel": 960, "programado": 300, "executado": 270,
                     "pct_programado": 0.3125, "aderencia": 0.9, "n_ordens": 1} for d in range(7)}}
    resumo = {d: {"programado": 300, "disponivel": 960, "pct_programado": 0.3125,
                  "programado_ant": 240, "delta_pct": 0.25} for d in range(7)}
    return {"semana": SEMANA, "eh_futura": False, "coletado_em": datetime(2026, 9, 18, 6, 0),
            "disciplinas": {"Mecanica": {"tecnicos": mec or [], "capacidade": cap},
                            "Eletrica": {"tecnicos": ele or [], "capacidade": {}}},
            "a_planejar": {"Mecanica": a_planejar or [], "Eletrica": []},
            "terceiros": terceiros or [], "pendentes_classificacao": pendentes or [],
            "resumo_dias": resumo, "meta_pct": 0.85}


def _montar(quadro, atrib=None, replan=None, carry=None):
    return px.montar_planilha(quadro, atrib or {}, replan or [], carry or {}, "rodolfo",
                              EXPORTADO_EM)


def _aba(modelo, titulo):
    aba = next(a for a in modelo["abas"] if a["titulo"] == titulo)
    nomes = [c[0] for c in aba["colunas"]]
    return [dict(zip(nomes, linha)) for linha in aba["linhas"]]


# ---- nome do arquivo e abas -----------------------------------------------------------

def test_nome_do_arquivo_traz_a_semana():
    assert _montar(_quadro())["nome_arquivo"] == "plano-semanal_2026-W38.xlsx"


def test_todas_as_abas_presentes():
    titulos = [a["titulo"] for a in _montar(_quadro())["abas"]]
    assert titulos == ["Atribuições", "Ordens", "A planejar", "Resumo por dia",
                       "Capacidade por turno", "Replanejamentos", "Sobre"]


# ---- Atribuições -------------------------------------------------------------------------

def test_ordem_dividida_vira_uma_linha_por_card_com_realizado_rateado():
    """5h planejadas, 4h30 realizadas, divididas 2h (João/ter) + 3h (Maria/qua): o realizado
    rateado soma 4h30 — não 9h (realizado cheio repetido em cada linha)."""
    joao = _tec("f1", "João", {1: [_ordem("5018123", 300, 270, parcela=120, atrib_id="a1")]})
    maria = _tec("f2", "Maria", {2: [_ordem("5018123", 300, 270, parcela=180, atrib_id="a2")]})
    linhas = _aba(_montar(_quadro(mec=[joao, maria])), "Atribuições")
    assert len(linhas) == 2
    assert [l["Técnico"] for l in linhas] == ["João", "Maria"]       # ordenado por data
    assert [l["Planejado na parcela (h)"] for l in linhas] == [2.0, 3.0]
    assert round(sum(l["Realizado rateado (h)"] for l in linhas), 2) == 4.5
    assert linhas[0]["Data"] == hh_data.datas_da_semana(SEMANA)[1].date()
    assert linhas[0]["Dia"] == "Terça"
    assert linhas[0]["Ordem inteira?"] == "Não"


def test_ordem_inteira_usa_planejado_da_ordem():
    tec = _tec("f1", "João", {0: [_ordem("5019794", 400, 0, parcela=None)]})
    linha = _aba(_montar(_quadro(mec=[tec])), "Atribuições")[0]
    assert linha["Ordem inteira?"] == "Sim"
    assert linha["Planejado na parcela (h)"] == round(400 / 60, 2)


def test_quem_e_quando_vem_da_atribuicao():
    tec = _tec("f1", "João", {0: [_ordem("1", 60, parcela=60, atrib_id="abc")]})
    atrib = {"abc": {"por": "pcm.ana", "criado_em": datetime(2026, 9, 12, 8, 0),
                     "atualizado_em": datetime(2026, 9, 13, 9, 0)}}
    linha = _aba(_montar(_quadro(mec=[tec]), atrib=atrib), "Atribuições")[0]
    assert linha["Atribuído por"] == "pcm.ana"
    assert linha["Atribuído em"] == datetime(2026, 9, 12, 8, 0)
    assert linha["Última alteração"] == datetime(2026, 9, 13, 9, 0)


def test_situacao_da_ordem_fechada_e_aguardando():
    fechada = _ordem("2", 60, 60, parcela=60, _entregue=True)
    aguardando = {"ordem": "3", "descricao": "(aguardando fechamento no SAP)",
                  "_aguardando_fechamento": True, "_atrib_id": "x", "parcela_min": 30}
    tec = _tec("f1", "João", {0: [fechada, aguardando], 1: [_ordem("4", 60, parcela=60)]})
    sit = {l["Ordem"]: l["Situação no SAP"] for l in _aba(_montar(_quadro(mec=[tec])), "Atribuições")}
    assert sit == {"2": "Fechada no SAP", "3": "Saiu do backlog (aguardando fechamento)",
                   "4": "Aberta"}


# ---- Ordens ------------------------------------------------------------------------------

def test_ordem_agregada_uma_linha_com_realizado_cheio():
    joao = _tec("f1", "João", {1: [_ordem("5018123", 300, 270, parcela=120, atrib_id="a1")]})
    maria = _tec("f2", "Maria", {2: [_ordem("5018123", 300, 270, parcela=180, atrib_id="a2")]})
    linhas = _aba(_montar(_quadro(mec=[joao, maria]), carry={"5018123": 2}), "Ordens")
    assert len(linhas) == 1
    o = linhas[0]
    assert o["Planejado da ordem (h)"] == 5.0
    assert o["Planejado na semana (h)"] == 5.0
    assert o["Realizado total da ordem (h)"] == 4.5
    assert o["Aderência"] == 0.9
    assert o["Técnicos"] == "João, Maria"
    assert o["Dias"] == "Terça, Quarta"
    assert o["Semanas anteriores planejada"] == 2
    assert o["Executores no SAP"] == "JOAO"          # sem de/para → o próprio código


def test_executores_traduzidos_pelo_de_para():
    tec = _tec("f1", "João", {0: [_ordem("7", 60, executores=["89990033", "X"])]})
    m = px.montar_planilha(_quadro(mec=[tec]), {}, [], {}, "r", EXPORTADO_EM,
                           {"89990033": "Pessoa de exemplo 17"})
    assert _aba(m, "Ordens")[0]["Executores no SAP"] == "Pessoa de exemplo 17, X"


# ---- A planejar --------------------------------------------------------------------------

def test_fila_a_planejar_traz_todas_as_filas_e_o_que_falta():
    parcial = _ordem("10", 300, parcela=None, _restante_min=120, _idade_dias=40)
    terc = _ordem("11", 60, _idade_dias=3)
    pend = _ordem("12", 0, _idade_dias=90)
    linhas = _aba(_montar(_quadro(a_planejar=[parcial], terceiros=[terc], pendentes=[pend])),
                  "A planejar")
    assert [(l["Fila"], l["Ordem"]) for l in linhas] == [
        ("Mecânica", "10"), ("Terceiros", "11"), ("Pendente de classificação", "12")]
    assert linhas[0]["Falta planejar (h)"] == 2.0
    assert linhas[1]["Falta planejar (h)"] == 1.0      # sem restante calculado = planejado todo
    assert linhas[0]["Idade (dias)"] == 40


# ---- Resumo / Capacidade / Replanejamentos ---------------------------------------------

def test_resumo_por_dia_e_capacidade():
    m = _montar(_quadro())
    dias = _aba(m, "Resumo por dia")
    assert len(dias) == 7
    assert dias[0]["Planejado (h)"] == 5.0 and dias[0]["Disponível (h)"] == 16.0
    assert dias[0]["Ocupação"] == 0.3125 and dias[0]["Variação vs semana anterior"] == 0.25
    cap = _aba(m, "Capacidade por turno")
    assert len(cap) == 7 and cap[0]["Quadro"] == "Mecânica" and cap[0]["Turno"] == "A"
    assert cap[0]["Realizado rateado (h)"] == 4.5


def test_replanejamentos_em_ordem_cronologica_com_rotulo():
    rep = [{"em": datetime(2026, 9, 16), "por": "b", "ordem": "2", "tipo": "remover",
            "categoria": "Falta de peça", "texto": ""},
           {"em": datetime(2026, 9, 15), "por": "a", "ordem": "1", "tipo": "remover",
            "categoria": "Outro", "texto": "parada da linha"}]
    linhas = _aba(_montar(_quadro(), replan=rep), "Replanejamentos")
    assert [l["Ordem"] for l in linhas] == ["1", "2"]
    assert linhas[0]["Tipo"] == "Remoção do plano"
    assert linhas[0]["Detalhe do motivo"] == "parada da linha"


# ---- render ------------------------------------------------------------------------------

def test_gerar_xlsx_relido_bate_com_o_modelo():
    joao = _tec("f1", "João", {1: [_ordem("5018123", 300, 270, parcela=120, atrib_id="a1")]})
    modelo = _montar(_quadro(mec=[joao]))
    wb = load_workbook(io.BytesIO(px.gerar_xlsx(modelo)))
    assert wb.sheetnames == [a["titulo"] for a in modelo["abas"]]
    ws = wb["Atribuições"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
    cab = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=cab.index("Ordem") + 1).value == "5018123"
    assert ws.cell(row=2, column=cab.index("Planejado na parcela (h)") + 1).value == 2.0
    assert wb["Sobre"]["B1"].value == SEMANA


def test_gerar_xlsx_semana_vazia_ainda_gera_arquivo():
    wb = load_workbook(io.BytesIO(px.gerar_xlsx(_montar(_quadro()))))
    assert wb["Atribuições"].max_row == 1      # só o cabeçalho
