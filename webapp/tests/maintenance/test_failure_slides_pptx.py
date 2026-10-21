"""Testes do gerador do material mensal em PowerPoint (S3) e do desenho da árvore.

Gera o arquivo em memória e inspeciona a apresentação resultante — sem abrir o
PowerPoint e sem tocar no banco.
"""

import io
import unittest

from pptx import Presentation

from src.utils.failure_slides import montar_mes
from src.utils import failure_slides_pptx as fsp
from src.utils.failure_slides_pptx import gerar_analise, gerar_pptx, nome_arquivo
from src.utils.failure_tree_image import desenhar


def _no(id_, parent, tipo, titulo="", ordem=0, minutos=None, raiz=False):
    return {"id": id_, "parent_id": parent, "tipo": tipo, "titulo": titulo,
            "sort_order": ordem, "duration_min": minutos, "is_root_cause": raiz}


def _analise_trabalhada(**kw):
    base = {
        "id": "6a59204d9bd97f1b60a3c0fe",
        "parada_chave": "LCT002_202607_201_abc12345",
        "titulo": "Falha no acionamento da tesoura",
        "status": "aberta",
        "responsavel": "Thiago Nicaretta",
        "parada": {"equipamento": "TRANS003", "equipamento_nome": "LCT-2,5",
                   "causa": "201", "duracao_min": 672, "n_apontamentos": 2,
                   "ordens": [5020308]},
        "nos": [_no("raiz", None, "ROOT", "Parada da LCT-2,5", minutos=672),
                _no("s1", "raiz", "STEP", "Baixa resistência de isolação", minutos=672),
                _no("c1", "s1", "CAUSE", "Umidade no enrolamento", minutos=400, raiz=True),
                _no("c2", "s1", "CAUSE", "Sem manutenção preventiva", minutos=272)],
        "problema": {"descricao": "Motor da tesoura com baixa resistência de isolação.",
                     "area": "Linha de corte", "observacoes": ""},
        "contencao": {"acoes": "Substituição do motor."},
        "actions": [{"node_id": "c1", "descricao": "Instalar resistência de aquecimento",
                     "responsavel": "Thiago", "prazo": "2026-09-15", "status": 2}],
        "sustainability": {"licoes": "Megôhmetro entrou no plano semanal."},
        "closing": {"responsavel": "", "data": None},
        "fotos": [],
    }
    base.update(kw)
    return base


def _analise_nao_iniciada():
    a = _analise_trabalhada()
    a.update({
        "parada_chave": "PRENS001_202607_201_def67890",
        "titulo": "GUILHOTINA RISCANDO MATERIAL",
        "nos": [_no("raiz", None, "ROOT", "Parada"), _no("s1", "raiz", "STEP", "relato")],
        "problema": {"descricao": "", "area": "", "observacoes": ""},
        "contencao": {"acoes": ""},
        "actions": [],
        "sustainability": {"licoes": ""},
        "parada": {"equipamento_nome": "PRENSA-01", "duracao_min": 188},
    })
    return a


def _textos(slide) -> str:
    return " ".join(f.text_frame.text for f in slide.shapes if f.has_text_frame)


class TestGerarPptx(unittest.TestCase):
    def setUp(self):
        material = montar_mes([_analise_trabalhada(), _analise_nao_iniciada()], 2026, 7)
        self.prs = Presentation(io.BytesIO(gerar_pptx(material)))
        self.slides = list(self.prs.slides)

    def test_capa_mais_um_slide_por_analise_mais_a_lista(self):
        self.assertEqual(len(self.slides), 3)

    def test_capa_traz_o_mes_por_extenso(self):
        self.assertIn("julho de 2026", _textos(self.slides[0]))

    def test_slide_da_analise_traz_equipamento_duracao_e_secoes(self):
        t = _textos(self.slides[1])
        for esperado in ("LCT-2,5", "11 h 12 min", "1. PROBLEMA", "2. CONTENÇÃO",
                         "4. PLANO DE AÇÃO", "3. ÁRVORE DE CAUSAS"):
            self.assertIn(esperado, t)

    def test_texto_das_secoes_e_texto_de_verdade_no_slide(self):
        t = _textos(self.slides[1])
        self.assertIn("baixa resistência de isolação", t)
        self.assertIn("Megôhmetro entrou no plano semanal.", t)

    def test_acao_mostra_causa_do_balao_e_status_por_extenso(self):
        t = _textos(self.slides[1])
        self.assertIn("Umidade no enrolamento", t)
        self.assertIn("Em andamento", t)

    def test_arvore_entra_como_imagem(self):
        imagens = [s for s in self.slides[1].shapes if s.shape_type == 13]
        self.assertGreaterEqual(len(imagens), 1)

    def test_nao_iniciadas_viram_uma_tabela_no_fim(self):
        ultimo = self.slides[-1]
        self.assertIn("sem tratamento", _textos(ultimo))
        tabelas = [s for s in ultimo.shapes if s.has_table]
        self.assertEqual(len(tabelas), 1)
        self.assertEqual(len(tabelas[0].table.rows), 2)       # cabeçalho + 1 pendente

    def test_formato_16_por_9(self):
        self.assertAlmostEqual(self.prs.slide_width / self.prs.slide_height, 16 / 9, places=2)

    def test_nada_estoura_a_borda_do_slide(self):
        L, A = self.prs.slide_width, self.prs.slide_height
        fora = [s for slide in self.slides for s in slide.shapes
                if s.left is not None and (s.left < 0 or s.top < 0
                                           or s.left + s.width > L + 1000
                                           or s.top + s.height > A + 1000)]
        self.assertEqual(fora, [])


class TestCasosDeBorda(unittest.TestCase):
    def test_mes_sem_analise_gera_so_a_capa(self):
        prs = Presentation(io.BytesIO(gerar_pptx(montar_mes([], 2026, 8))))
        self.assertEqual(len(list(prs.slides)), 1)

    def test_mes_so_com_nao_iniciadas_nao_tem_slide_de_analise(self):
        material = montar_mes([_analise_nao_iniciada()], 2026, 7)
        prs = Presentation(io.BytesIO(gerar_pptx(material)))
        self.assertEqual(len(list(prs.slides)), 2)           # capa + lista

    def test_analise_sem_acao_diz_que_nao_ha_acao(self):
        a = _analise_trabalhada(actions=[])
        prs = Presentation(io.BytesIO(gerar_pptx(montar_mes([a], 2026, 7))))
        self.assertIn("Nenhuma ação registrada", _textos(list(prs.slides)[1]))

    def test_texto_gigante_e_cortado_sem_estourar(self):
        a = _analise_trabalhada()
        a["problema"] = dict(a["problema"], descricao="x" * 5000)
        prs = Presentation(io.BytesIO(gerar_pptx(montar_mes([a], 2026, 7))))
        texto = _textos(list(prs.slides)[1])
        self.assertIn("…", texto)
        self.assertLess(len(texto), 5000)

    def test_foto_inexistente_nao_derruba_a_geracao(self):
        a = _analise_trabalhada(fotos=[{"arquivo": "sumiu.jpg", "legenda": "x"}])
        prs = Presentation(io.BytesIO(gerar_pptx(montar_mes([a], 2026, 7))))
        self.assertIn("Sem fotos anexadas", _textos(list(prs.slides)[1]))

    def test_nome_do_arquivo_ordena_sozinho(self):
        self.assertEqual(nome_arquivo(montar_mes([], 2026, 7)),
                         "analises-falha-2026-07.pptx")


class TestDesenhoDaArvore(unittest.TestCase):
    def test_devolve_png(self):
        material = montar_mes([_analise_trabalhada()], 2026, 7)
        png = desenhar(material["slides"][0]["arvore"])
        self.assertTrue(png.startswith(b"\x89PNG"))

    def test_arvore_vazia_nao_desenha(self):
        self.assertIsNone(desenhar([]))

    def test_titulo_gigante_nao_quebra_o_desenho(self):
        arvore = [{"id": "r", "parent_id": None, "nivel": 0, "ordem": 0,
                   "titulo": "palavra " * 60, "tipo": "ROOT", "duration_min": 10,
                   "causa_raiz": False, "n_filhos": 0}]
        self.assertTrue(desenhar(arvore).startswith(b"\x89PNG"))


if __name__ == "__main__":
    unittest.main()


def _png_bytes(tam=(400, 120)) -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", tam, (255, 255, 255)).save(buf, format="PNG")
    return buf.getvalue()


_SVG_QUADRO = b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 120"><rect width="400" height="120"/></svg>'


def _partes_svg(dados: bytes) -> list:
    import zipfile
    return [n for n in zipfile.ZipFile(io.BytesIO(dados)).namelist() if n.endswith(".svg")]


def _tem_svg_blip(slide) -> bool:
    return any("svgBlip" in etree_tostring(sh._element) for sh in slide.shapes
               if sh.shape_type == 13)          # 13 = PICTURE


def etree_tostring(el) -> str:
    from lxml import etree
    return etree.tostring(el).decode()


class TestDiagramaDoQuadro(unittest.TestCase):
    """IM-AF3-08 — o painel 3 usa o desenho do próprio quadro (SVG + PNG de reserva)
    quando há imagem válida guardada; senão, o redesenho do servidor."""

    def _gerar(self, diagrama):
        from unittest import mock
        with mock.patch("src.utils.failure_slides.imagem_diagrama", return_value=diagrama):
            return gerar_analise(_analise_trabalhada())[0]

    def test_com_imagem_guardada_o_slide_leva_o_svg(self):
        dados = self._gerar({"svg": _SVG_QUADRO, "png": _png_bytes()})
        self.assertEqual(_partes_svg(dados), ["ppt/media/diagrama1.svg"])
        prs = Presentation(io.BytesIO(dados))
        self.assertTrue(_tem_svg_blip(prs.slides[0]))
        import zipfile
        tipos = zipfile.ZipFile(io.BytesIO(dados)).read("[Content_Types].xml").decode()
        self.assertIn('ContentType="image/svg+xml"', tipos)

    def test_sem_imagem_guardada_usa_o_redesenho(self):
        dados = self._gerar(None)
        self.assertEqual(_partes_svg(dados), [])
        prs = Presentation(io.BytesIO(dados))
        self.assertFalse(_tem_svg_blip(prs.slides[0]))
        self.assertTrue(any(sh.shape_type == 13 for sh in prs.slides[0].shapes))

    def test_imagem_corrompida_cai_no_redesenho_sem_derrubar(self):
        dados = self._gerar({"svg": _SVG_QUADRO, "png": b"nao-e-png"})
        prs = Presentation(io.BytesIO(dados))
        self.assertEqual(len(prs.slides), 1)
        self.assertFalse(_tem_svg_blip(prs.slides[0]))

    def test_imagem_fica_dentro_do_painel_da_arvore(self):
        dados = self._gerar({"svg": _SVG_QUADRO, "png": _png_bytes((4000, 300))})
        slide = Presentation(io.BytesIO(dados)).slides[0]
        pic = next(sh for sh in slide.shapes if sh.shape_type == 13 and "svgBlip" in etree_tostring(sh._element))
        x, y = fsp.LARGA_X + 0.3, 3.7
        self.assertGreaterEqual(pic.left, fsp._cm(x) - 1)
        self.assertLessEqual(pic.left + pic.width, fsp._cm(x + fsp.LARGA_W - 0.6) + 1)
        self.assertGreaterEqual(pic.top, fsp._cm(y) - 1)
        self.assertLessEqual(pic.top + pic.height, fsp._cm(y + 8.9) + 1)

    def test_material_mensal_com_duas_analises_tem_dois_svgs(self):
        from unittest import mock
        a1 = _analise_trabalhada()
        a2 = _analise_trabalhada(parada_chave="LCT002_202607_201_zzz99999", id="6a59204d9bd97f1b60a3c0ff")
        with mock.patch("src.utils.failure_slides.imagem_diagrama",
                        return_value={"svg": _SVG_QUADRO, "png": _png_bytes()}):
            dados = gerar_pptx(montar_mes([a1, a2], 2026, 7))
        self.assertEqual(sorted(_partes_svg(dados)), ["ppt/media/diagrama1.svg", "ppt/media/diagrama2.svg"])


class TestExportarUmaAnalise(unittest.TestCase):
    """S5 — "Exportar esta análise": .pptx de um slide só, sem capa nem lista do fim."""

    def test_gera_um_unico_slide(self):
        dados, _nome = gerar_analise(_analise_trabalhada())
        prs = Presentation(io.BytesIO(dados))
        self.assertEqual(len(prs.slides), 1)

    def test_slide_traz_o_conteudo_da_analise(self):
        dados, _ = gerar_analise(_analise_trabalhada())
        prs = Presentation(io.BytesIO(dados))
        texto = _textos(prs.slides[0])
        self.assertIn("LCT-2,5", texto)
        self.assertIn("baixa resistência de isolação", texto.lower())

    def test_nome_do_arquivo_traz_equipamento_e_mes(self):
        _dados, nome = gerar_analise(_analise_trabalhada())
        self.assertEqual(nome, "analise-falha-LCT002-2026-07.pptx")

    def test_chave_em_formato_antigo_nao_quebra_o_nome(self):
        a = _analise_trabalhada(parada_chave="LONGI002_202606220911_201")
        _dados, nome = gerar_analise(a)
        self.assertEqual(nome, "analise-falha-LONGI002.pptx")

    def test_analise_nao_iniciada_ainda_gera_o_slide(self):
        # O usuário pediu ESTA análise: exporta mesmo vazia, em vez de recusar.
        dados, nome = gerar_analise(_analise_nao_iniciada())
        prs = Presentation(io.BytesIO(dados))
        self.assertEqual(len(prs.slides), 1)
        self.assertEqual(nome, "analise-falha-PRENS001-2026-07.pptx")

    def test_nao_tem_capa_nem_lista_do_fim(self):
        dados, _ = gerar_analise(_analise_trabalhada())
        prs = Presentation(io.BytesIO(dados))
        texto = _textos(prs.slides[0])
        self.assertNotIn("Análises de Falha —", texto)
        self.assertNotIn("sem tratamento", texto)


def _caixas_texto(slide):
    """(x, y, largura, altura) em EMU de cada caixa de texto do slide, com o texto."""
    caixas = []
    for f in slide.shapes:
        if f.has_text_frame and f.text_frame.text.strip():
            caixas.append((f.left, f.top, f.width, f.height, f.text_frame.text))
    return caixas


class TestSemSobreposicaoDeTexto(unittest.TestCase):
    """Regressão do defeito visto no material de julho/2026: com 3+ ações de texto
    longo, as linhas do Plano de Ação se atropelavam e vazavam para a seção de baixo,
    deixando o slide ilegível justamente no bloco que a reunião discute."""

    def _analise_com_acoes_longas(self, n=4):
        a = _analise_trabalhada()
        a["actions"] = [
            {"node_id": f"c{i}",
             "descricao": ("Revisão do plano de manutenção com inclusão da inspeção "
                           "desse item no preventivo de motores e treinamento dos "
                           "mantenedores envolvidos na rotina " * 2),
             "responsavel": "Pessoa de exemplo 09", "prazo": "2026-09-30", "status": 2}
            for i in range(n)
        ]
        a["nos"] = [_no("raiz", None, "ROOT", "Parada", minutos=672),
                    _no("s1", "raiz", "STEP", "Etapa", minutos=672)] + [
            _no(f"c{i}", "s1", "CAUSE",
                "Falta de estrategia de manutenção (inspeção) para detecção de anomalias",
                minutos=100, raiz=True) for i in range(n)]
        return a

    def _sobrepostas(self, slide):
        """Pares de caixas cujo retângulo se cruza (tolerância de 1 pt em EMU)."""
        tol = 12700
        caixas = _caixas_texto(slide)
        pares = []
        for i, (x1, y1, w1, h1, t1) in enumerate(caixas):
            for (x2, y2, w2, h2, t2) in caixas[i + 1:]:
                cruza_x = x1 + w1 - tol > x2 and x2 + w2 - tol > x1
                cruza_y = y1 + h1 - tol > y2 and y2 + h2 - tol > y1
                if cruza_x and cruza_y:
                    pares.append((t1[:40], t2[:40]))
        return pares

    def test_plano_de_acao_nao_sobrepoe(self):
        material = montar_mes([self._analise_com_acoes_longas(4)], 2026, 7)
        prs = Presentation(io.BytesIO(gerar_pptx(material)))
        pares = self._sobrepostas(prs.slides[1])
        self.assertEqual(pares, [], f"caixas sobrepostas: {pares}")

    def test_seis_acoes_nao_vazam_do_painel(self):
        material = montar_mes([self._analise_com_acoes_longas(6)], 2026, 7)
        prs = Presentation(io.BytesIO(gerar_pptx(material)))
        self.assertEqual(self._sobrepostas(prs.slides[1]), [])

    def test_uma_analise_real_nao_sobrepoe(self):
        dados, _ = gerar_analise(self._analise_com_acoes_longas(3))
        prs = Presentation(io.BytesIO(dados))
        self.assertEqual(self._sobrepostas(prs.slides[0]), [])


class TestSecao4MostraTudo(unittest.TestCase):
    """A seção 4 (Plano de Ação) é o compromisso que sai da reunião: nenhuma ação pode
    ficar de fora e nenhuma frase pode terminar em "…" enquanto houver como diminuir a
    letra. Regressão do que o Rodolfo apontou em 2026-08-06."""

    def _acoes(self, n, tamanho=180):
        return [{"node_id": f"c{i}",
                 "descricao": ("Revisão do plano de manutenção com inclusão da inspeção "
                               "desse item no preventivo de motores " * 3)[:tamanho],
                 "responsavel": "Pessoa de exemplo 09", "prazo": "2026-09-30", "status": 2}
                for i in range(n)]

    def _analise(self, n, tamanho=180):
        a = _analise_trabalhada()
        a["actions"] = self._acoes(n, tamanho)
        a["nos"] = [_no("raiz", None, "ROOT", "Parada", minutos=672),
                    _no("s1", "raiz", "STEP", "Etapa", minutos=672)] + [
            _no(f"c{i}", "s1", "CAUSE",
                f"Causa raiz numero {i} com um titulo razoavelmente longo de verdade",
                minutos=50, raiz=True) for i in range(n)]
        return a

    def _texto_slide(self, analise):
        dados, _ = gerar_analise(analise)
        prs = Presentation(io.BytesIO(dados))
        return _textos(prs.slides[0])

    def test_todas_as_acoes_aparecem(self):
        texto = self._texto_slide(self._analise(4))
        for i in range(4):
            self.assertIn(f"Causa raiz numero {i}", texto)
        self.assertNotIn("ações não exibidas", texto)

    def test_seis_acoes_curtas_cabem_todas(self):
        texto = self._texto_slide(self._analise(6, tamanho=90))
        for i in range(6):
            self.assertIn(f"Causa raiz numero {i}", texto)

    def test_descricao_nao_termina_em_reticencias(self):
        texto = self._texto_slide(self._analise(3))
        self.assertNotIn("…", texto)

    def test_responsavel_e_prazo_aparecem_por_acao(self):
        texto = self._texto_slide(self._analise(2))
        self.assertIn("Pessoa de exemplo 09", texto)
        self.assertIn("30/09/26", texto)     # prazo no formato curto
        self.assertIn("Em andamento", texto)


class TestIdentidadeVisualDoCliente(unittest.TestCase):
    """Todo slide leva a moldura da identidade do cliente: faixa vertical azul na borda
    direita e o quarto de anel laranja no canto inferior direito (slide-modelo que ele
    mandou em 2026-08-06)."""

    def _prs(self):
        material = montar_mes([_analise_trabalhada(), _analise_nao_iniciada()], 2026, 7)
        return Presentation(io.BytesIO(gerar_pptx(material)))

    def test_todos_os_slides_tem_a_faixa_azul(self):
        from pptx.util import Cm
        prs = self._prs()
        self.assertGreaterEqual(len(prs.slides), 3)      # capa + análise + lista
        for i, slide in enumerate(prs.slides, 1):
            faixas = [f for f in slide.shapes
                      if abs(f.left - Cm(fsp.BARRA_X)) < Cm(0.05)
                      and f.height >= prs.slide_height - Cm(0.05)]
            self.assertTrue(faixas, f"slide {i} sem a faixa azul da identidade")

    def test_todos_os_slides_tem_o_arco_laranja(self):
        from pptx.util import Cm
        prs = self._prs()
        for i, slide in enumerate(prs.slides, 1):
            arcos = [f for f in slide.shapes
                     if f.shape_type == 13                      # imagem
                     and f.top + f.height >= prs.slide_height - Cm(0.05)
                     and f.left + f.width >= prs.slide_width - Cm(0.05)]
            self.assertTrue(arcos, f"slide {i} sem o arco laranja da identidade")

    def test_conteudo_nao_invade_a_faixa_lateral(self):
        from pptx.util import Cm
        prs = self._prs()
        limite = Cm(fsp.BARRA_X)
        for i, slide in enumerate(prs.slides, 1):
            for f in slide.shapes:
                if f.has_text_frame and f.text_frame.text.strip():
                    self.assertLessEqual(
                        f.left, limite,
                        f"slide {i}: texto {f.text_frame.text[:30]!r} começa dentro da faixa")

    def test_cores_sao_as_do_cliente(self):
        self.assertEqual(fsp.AZUL, (0, 83, 128))        # #005380
        self.assertEqual(fsp.LARANJA, (255, 55, 0))     # #ff3700


class TestVariasAcoesPorCausaNoSlide(unittest.TestCase):
    """Uma causa raiz pode ter várias ações. No slide, o nome da causa é escrito
    UMA vez por grupo — repetir o título em negrito a cada ação comia tanta altura
    que ações inteiras ficavam de fora ("+N ações não exibidas")."""

    @staticmethod
    def _acao(node_id, descricao, responsavel="", prazo=None, causa="Correia folgada"):
        return {"causa": causa, "descricao": descricao, "responsavel": responsavel,
                "prazo": prazo, "status": 1, "status_texto": "Aprovado",
                "node_id": node_id}

    def test_causa_repetida_escrita_uma_vez(self):
        acoes = [self._acao("a", "Trocar a correia"),
                 self._acao("a", "Revisar o tensor"),
                 self._acao("b", "Ajustar rota", causa="Mancal sem óleo")]
        itens, _ = fsp._medir_acoes(acoes, 11.0, 0)
        # 1ª de cada causa escreve o título; a 2ª da mesma causa não.
        self.assertGreater(itens[0]["n_causa"], 0)
        self.assertEqual(itens[1]["n_causa"], 0)
        self.assertGreater(itens[2]["n_causa"], 0)

    def test_omitir_a_causa_repetida_encurta_o_bloco(self):
        # É o ponto: sem isso, as ações extras empurram outras para fora do slide.
        uma = [self._acao("a", "Trocar a correia")]
        duas = uma + [self._acao("a", "Revisar o tensor")]
        _, alt_uma = fsp._medir_acoes(uma, 11.0, 0)
        _, alt_duas = fsp._medir_acoes(duas, 11.0, 0)
        # A 2ª ação custa menos que a 1ª, justamente por não repetir o título.
        self.assertLess(alt_duas - alt_uma, alt_uma)

    def test_causa_igual_nao_consecutiva_volta_a_ser_escrita(self):
        # A ordem vem agrupada por causa; se por algum motivo não vier, o título
        # reaparece — melhor repetir do que deixar ação órfã de causa.
        acoes = [self._acao("a", "Trocar"), self._acao("b", "Ajustar", causa="Outra"),
                 self._acao("a", "Revisar")]
        itens, _ = fsp._medir_acoes(acoes, 11.0, 0)
        self.assertGreater(itens[2]["n_causa"], 0)

    def test_varios_responsaveis_cabem_sem_invadir_a_acao_seguinte(self):
        # O rodapé (quem · até quando) pode passar de uma linha com vários nomes;
        # medir só uma linha fazia o texto invadir o bloco de baixo.
        curto = [self._acao("a", "Trocar", responsavel="João", prazo="2026-08-10")]
        longo = [self._acao("a", "Trocar", prazo="2026-08-10",
                            responsavel="Pessoa de exemplo 09, Pessoa de exemplo 20, "
                                        "Rodolfo J. Silva, Joaquim Soares, José da Silva")]
        i_curto, _ = fsp._medir_acoes(curto, 11.0, 0)
        i_longo, _ = fsp._medir_acoes(longo, 11.0, 0)
        self.assertEqual(i_curto[0]["n_meta"], 1)
        self.assertGreater(i_longo[0]["n_meta"], 1)
        self.assertGreater(i_longo[0]["altura"], i_curto[0]["altura"])

    def test_slide_sai_sem_erro_com_varias_acoes_e_responsaveis(self):
        analise = {
            "titulo": "Eixo quebrou", "parada": {"equipamento_nome": "LCT-08",
                                                 "duracao_min": 155},
            "nos": [], "actions": [
                {"id": "1", "node_id": "a", "descricao": "Trocar a correia",
                 "responsavel": ["Pessoa de exemplo 09", "Pessoa de exemplo 20"],
                 "prazo": "2026-08-10", "status": 1},
                {"id": "2", "node_id": "a", "descricao": "Revisar o tensor",
                 "responsavel": ["Rodolfo J. Silva"], "prazo": "2026-09-15",
                 "status": 2, "adicionada": True},
            ],
        }
        dados, _ = gerar_analise(analise)
        self.assertGreater(len(dados), 0)
        prs = Presentation(io.BytesIO(dados))
        textos = " ".join(sh.text_frame.text for sl in prs.slides
                          for sh in sl.shapes if sh.has_text_frame)
        self.assertIn("Pessoa de exemplo 09, Pessoa de exemplo 20", textos)
        self.assertIn("Revisar o tensor", textos)
        # O título da causa não pode sair duplicado.
        self.assertEqual(textos.count("Correia folgada"), 0)  # sem nós, causa vazia
