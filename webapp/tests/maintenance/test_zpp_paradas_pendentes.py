"""Testes das paradas pendentes de consolidação na NT0001 (KPIReport-v2 BR-14 / SP-26/28/29).

Sem Mongo real: `_get_coll` e `fetch_top_breakdowns_all` são trocados por fakes. Casos com
dados do estudo de 07/10/2026 (LONGI001, ordem 10624555, quebra do redutor, 345 min).
"""
from __future__ import annotations

from datetime import datetime

import pytest

from src.utils import kpi_report_data as krd
from src.utils import zpp_paradas_pendentes as zpend
from src.utils import zpp_paradas_scheduler as zsched

_INI, _FIM = datetime(2026, 10, 6), datetime(2026, 10, 7)
_AGORA = datetime(2026, 10, 7, 7, 30)


def _bate(doc, q):
    """Match mínimo: igualdade, $in, $nin, $ne, $gte/$lt."""
    for k, cond in q.items():
        v = doc.get(k)
        if isinstance(cond, dict):
            if "$in" in cond and v not in cond["$in"]:
                return False
            if "$nin" in cond and v in cond["$nin"]:
                return False
            if "$ne" in cond and v == cond["$ne"]:
                return False
            if "$gte" in cond and not (v is not None and v >= cond["$gte"]):
                return False
            if "$lt" in cond and not (v is not None and v < cond["$lt"]):
                return False
        elif v != cond:
            return False
    return True


class _Coll:
    def __init__(self, docs):
        self.docs = docs

    def find(self, q, proj=None):
        return [dict(d) for d in self.docs if _bate(d, q)]

    def find_one(self, q):
        return next((dict(d) for d in self.docs if _bate(d, q)), None)


def _retrato(*paradas, coletado_em=_AGORA):
    return [{"_id": "meta", "coletado_em": coletado_em, "total": len(paradas)}] + list(paradas)


def _p(linha, ordem, inicio, dur, causa="201", texto="t"):
    return {"linha": linha, "ordem": ordem, "inicio": inicio, "duracao_min": dur, "causa": causa,
            "texto_operador": texto, "descricao_motivo": "d", "coleta_id": "c"}


def _of(centro, ordem, inicio, causa="201"):
    return {"centro_de_trabalho": centro, "ordem": ordem,
            "inicio_execucao": inicio.replace(hour=0, minute=0, second=0),
            "inicio_real_hora": inicio.strftime("%H:%M:%S"), "causa_do_desvio": causa}


@pytest.fixture
def mongo(monkeypatch):
    dbs = {"ZPP_Paradas_Lancadas": _Coll([]), "ZPP_Paradas": _Coll([])}
    monkeypatch.setattr(zpend, "_get_coll", lambda name: dbs[name])
    return dbs


# ---------------------------------------------------------------- módulo de pendentes

def test_pendente_e_a_que_falta_na_oficial_pela_chave_exata(mongo):
    longi = datetime(2026, 10, 6, 6, 15)
    prens = datetime(2026, 10, 6, 9, 0, 1)
    mongo["ZPP_Paradas_Lancadas"].docs = _retrato(
        _p("LONGI001", 10624555, longi, 345.0, texto="Quebra redutor"),
        _p("PRENS001", 10626620, prens, 40.0),
    )
    # oficial guarda a ordem como string com zeros e a hora em texto
    mongo["ZPP_Paradas"].docs = [_of("PRENS001", "000010626620", prens)]
    pend = zpend.buscar_pendentes(_INI, _FIM)
    assert [(p["equipamento"], p["duration_min"], p["pendente"]) for p in pend] == [("LONGI001", 345.0, True)]
    assert pend[0]["descricao"] == "Quebra redutor"


def test_mesma_ordem_varias_paradas_casa_cada_uma_pelo_inicio(mongo):
    a, b = datetime(2026, 10, 6, 6, 15), datetime(2026, 10, 6, 13, 0, 10)
    mongo["ZPP_Paradas_Lancadas"].docs = _retrato(_p("LONGI001", 1, a, 345.0), _p("LONGI001", 1, b, 120.0))
    mongo["ZPP_Paradas"].docs = [_of("LONGI001", 1, a)]
    assert [p["inicio_execucao"] for p in zpend.buscar_pendentes(_INI, _FIM)] == [b]


def test_ja_na_oficial_com_outro_codigo_nao_e_pendente(mongo):
    """Caso 401 × S401 do estudo: se a parada existe na NT0001, não é pendente, qualquer o código."""
    t = datetime(2026, 10, 6, 13, 1)
    mongo["ZPP_Paradas_Lancadas"].docs = _retrato(_p("SOLDA001", 10631108, t, 64.0, causa="201"))
    mongo["ZPP_Paradas"].docs = [_of("SOLDA001", 10631108, t, causa="S401")]
    assert zpend.buscar_pendentes(_INI, _FIM) == []


def test_so_avaria_dentro_da_janela_e_fora_dos_excluidos(mongo):
    mongo["ZPP_Paradas_Lancadas"].docs = _retrato(
        _p("SOLDA001", 1, datetime(2026, 10, 6, 8), 10.0, causa="304"),      # set-up: não é avaria
        _p("PRENS001", 2, datetime(2026, 10, 7, 0, 0), 30.0),                # fora: fim é exclusivo
        _p("DECAP001", 3, datetime(2026, 10, 6, 9), 50.0),                   # equipamento zerado
        _p("PRENS002", 4, datetime(2026, 10, 6, 10), 20.0, causa="S202"),    # avaria com S
    )
    assert [p["equipamento"] for p in zpend.buscar_pendentes(_INI, _FIM)] == ["PRENS002"]


def test_falha_no_mongo_vira_lista_vazia(monkeypatch):
    def _boom(name):
        raise RuntimeError("mongo fora")
    monkeypatch.setattr(zpend, "_get_coll", _boom)
    assert zpend.buscar_pendentes(_INI, _FIM) == []
    assert zpend.retrato_valido(_AGORA) == (False, None)


def test_retrato_so_vale_se_coletado_hoje(mongo):
    mongo["ZPP_Paradas_Lancadas"].docs = _retrato(coletado_em=datetime(2026, 10, 7, 6, 31))
    assert zpend.retrato_valido(_AGORA) == (True, datetime(2026, 10, 7, 6, 31))
    mongo["ZPP_Paradas_Lancadas"].docs = _retrato(coletado_em=datetime(2026, 10, 6, 10, 2))
    ok, quando = zpend.retrato_valido(_AGORA)
    assert not ok and "06/10 10:02" in zpend.aviso_checagem(quando)
    mongo["ZPP_Paradas_Lancadas"].docs = []
    assert zpend.retrato_valido(_AGORA) == (False, None)
    assert "nenhuma coleta" in zpend.aviso_checagem(None)


# ---------------------------------------------------------------- mistura nas listas

def _oficial(eq, inicio, dur):
    return {"equipamento": eq, "inicio_execucao": inicio, "duration_min": dur,
            "causa_do_desvio": "201", "descricao": "of"}


_OFICIAIS = [_oficial(f"EQ{i}", datetime(2026, 10, 6, i), float(100 - i * 10)) for i in range(1, 7)]  # 90..40


@pytest.fixture
def lista(monkeypatch):
    chamadas = []

    def _fetch(ini, fim, top_n=5, **kw):
        chamadas.append(top_n)
        return _OFICIAIS if top_n is None else _OFICIAIS[:top_n]
    monkeypatch.setattr(krd, "fetch_top_breakdowns_all", _fetch)
    estado = {"ok": True, "coletado_em": _AGORA, "pend": []}
    monkeypatch.setattr(zpend, "retrato_valido", lambda agora: (estado["ok"], estado["coletado_em"]))
    monkeypatch.setattr(zpend, "buscar_pendentes", lambda ini, fim: list(estado["pend"]))
    return estado, chamadas


def test_sem_pendente_lista_identica_a_de_antes(lista):
    estado, chamadas = lista
    r = krd._build_top5_paradas(_INI, _FIM, {}, top_n=5, agora=_AGORA)
    assert [p["equipamento"] for p in r["paradas"]] == ["EQ1", "EQ2", "EQ3", "EQ4", "EQ5"]
    assert chamadas == [5]                       # mesmo caminho de sempre (corte no Mongo)
    assert r["tem_pendente"] is False and r["aviso_checagem"] == ""
    assert all(p["pendente"] is False for p in r["paradas"])


def test_pendente_disputa_posicao_e_empurra_a_quinta(lista):
    estado, chamadas = lista
    estado["pend"] = [{"equipamento": "LONGI001", "inicio_execucao": datetime(2026, 10, 6, 6, 15),
                       "duration_min": 85.0, "causa_do_desvio": "201", "descricao": "Quebra",
                       "pendente": True}]
    r = krd._build_top5_paradas(_INI, _FIM, {"LONGI001": "Longitudinal 1"}, top_n=5, agora=_AGORA)
    assert [p["equipamento"] for p in r["paradas"]] == ["EQ1", "Longitudinal 1", "EQ2", "EQ3", "EQ4"]
    assert [p["posicao"] for p in r["paradas"]] == [1, 2, 3, 4, 5]
    assert r["paradas"][1]["pendente"] is True and r["paradas"][1]["duracao_min"] == "85 min"
    assert r["tem_pendente"] is True
    assert chamadas == [None]                    # oficiais sem corte, corte depois da mistura


def test_bloco5_sem_corte_lista_todas(lista):
    estado, _ = lista
    estado["pend"] = [{"equipamento": "X", "inicio_execucao": datetime(2026, 10, 6, 5),
                       "duration_min": 5.0, "causa_do_desvio": "202", "descricao": "", "pendente": True}]
    r = krd._build_top5_paradas(_INI, _FIM, {}, top_n=None, agora=_AGORA)
    assert len(r["paradas"]) == 7 and r["paradas"][-1]["equipamento"] == "X"


def test_retrato_vencido_mostra_aviso_e_so_oficiais(lista):
    estado, chamadas = lista
    estado.update(ok=False, coletado_em=datetime(2026, 10, 6, 10, 0))
    r = krd._build_top5_paradas(_INI, _FIM, {}, top_n=5, agora=_AGORA)
    assert len(r["paradas"]) == 5 and not r["tem_pendente"]
    assert "06/10 10:00" in r["aviso_checagem"]


def test_tela_v1_nao_consulta_pendentes(lista, monkeypatch):
    estado, chamadas = lista
    estado["pend"] = [{"equipamento": "X", "inicio_execucao": _INI, "duration_min": 999.0,
                       "causa_do_desvio": "201", "descricao": "", "pendente": True}]
    r = krd._build_top5_paradas(_INI, _FIM, {}, top_n=5, incluir_pendentes=False)
    assert "X" not in [p["equipamento"] for p in r["paradas"]] and r["aviso_checagem"] == ""


def test_erro_na_checagem_nao_derruba_a_lista(lista, monkeypatch):
    def _boom(agora):
        raise RuntimeError("x")
    monkeypatch.setattr(zpend, "retrato_valido", _boom)
    r = krd._build_top5_paradas(_INI, _FIM, {}, top_n=5, agora=_AGORA)
    assert len(r["paradas"]) == 5 and not r["tem_pendente"]


# ---------------------------------------------------------------- agenda

def test_parse_horas_default_e_invalidas():
    assert zsched.parse_horas("06:30,10:00") == [(6, 30), (10, 0)]
    assert zsched.parse_horas("") == [(6, 30), (10, 0)]
    assert zsched.parse_horas("25:00, 7:05") == [(7, 5)]


def test_tipo_fora_do_agendador_padrao():
    """DS-18 revisado: não entra em TIPOS_VALIDOS (home e /config/sap-scheduler intactas)."""
    from src.sap_scheduler.validators import TIPOS_VALIDOS
    assert zsched.TIPO not in TIPOS_VALIDOS


# ---------------------------------------------------------------- desenho (SP-27)

def _linhas(pendente_na=2):
    return [{"posicao": i, "equipamento": f"EQ{i}", "data_hora": "06/10/2026 06:15",
             "duracao_min": "10 min", "causa": "201", "descricao": "x", "pendente": i == pendente_na}
            for i in range(1, 4)]


def _textos(comp):
    """Todos os textos e classes de uma árvore Dash (para asserções simples)."""
    out = []
    def walk(c):
        if c is None:
            return
        if isinstance(c, (str, int, float)):
            out.append(str(c)); return
        if isinstance(c, (list, tuple)):
            for x in c:
                walk(x)
            return
        cls = getattr(c, "className", None)
        if cls:
            out.append(f"class:{cls}")
        walk(getattr(c, "children", None))
    walk(comp)
    return out


def test_tela_v2_selo_fundo_e_legenda_so_com_pendente():
    from src.utils import kpi_report_v2_layout as lay
    com = _textos(lay.build_bloco5_html({"paradas": _linhas(2), "vazio": False}, "pt"))
    assert zpend.SELO in com and zpend.LEGENDA in com
    assert sum(1 for t in com if "kpi-v2-row-pendente" in t) == 1
    sem = _textos(lay.build_bloco5_html({"paradas": _linhas(None), "vazio": False}, "pt"))
    assert zpend.SELO not in sem and zpend.LEGENDA not in sem
    assert not any("kpi-v2-row-pendente" in t for t in sem)


def test_tela_v2_aviso_de_checagem():
    from src.utils import kpi_report_v2_layout as lay
    txt = _textos(lay.build_bloco2_html({"paradas": _linhas(None), "vazio": False,
                                         "aviso_checagem": "Checagem indisponível"}, "pt"))
    assert "Checagem indisponível" in txt


def test_pdf_com_pendente_tem_selo_e_legenda():
    from src.utils import kpi_report_v2_pdf as pdf
    styles = pdf._styles()
    fl = pdf._bloco_paradas_flowables(5, {"paradas": _linhas(2), "vazio": False}, "pt",
                                      "block_5_title", styles, top_n=None)
    textos = " ".join(getattr(f, "text", "") for f in fl)
    assert "PENDENTE" in textos and "não considerada no cálculo" in textos


def test_docx_marca_causa_e_insere_legenda():
    pytest.importorskip("docx")
    from src.utils import kpi_report_builder as kb
    dados = {"bloco2": {"paradas": _linhas(None)}, "bloco5": {"paradas": _linhas(2), "aviso_checagem": ""}}
    kb._marcar_pendentes_docx(dados)
    assert dados["bloco5"]["paradas"][1]["causa"] == "PENDENTE · 201"
    assert dados["bloco5"]["paradas"][0]["causa"] == "201"
    from docx import Document
    from io import BytesIO
    doc = Document()
    for _ in range(2):
        t = doc.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "#", "Equipamento"
        doc.add_paragraph("depois")
    buf = BytesIO(); doc.save(buf)
    out = Document(BytesIO(kb._inserir_legenda_pendentes_docx(buf.getvalue(), dados)))
    textos = [p.text for p in out.paragraphs]
    assert sum("não considerada no cálculo" in t for t in textos) == 1
    sem = {"bloco2": {"paradas": _linhas(None)}, "bloco5": {"paradas": _linhas(None)}}
    assert kb._inserir_legenda_pendentes_docx(b"x", sem) == b"x"


def test_aba_relatorio_tabela_compartilhada_selo_legenda_e_aviso():
    """Aba Relatório do indicators-v2 desenha com `_tabela_top_paradas` (mesma da V1)."""
    from src.callbacks_registers.kpi_report_screen_callbacks import _tabela_top_paradas
    com = _textos(_tabela_top_paradas(_linhas(2), False, "vazio", aviso="Checagem indisponível"))
    assert zpend.SELO in com and zpend.LEGENDA in com and "Checagem indisponível" in com
    assert sum(1 for t in com if "kpi-v2-row-pendente" in t) == 1
    # V1 (sem pendente, sem aviso): devolve a própria tabela, como antes
    import dash_bootstrap_components as dbc
    assert isinstance(_tabela_top_paradas(_linhas(None), False, "vazio"), dbc.Table)
    assert isinstance(_tabela_top_paradas([], True, "vazio"), dbc.Alert)
