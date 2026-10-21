"""Carregamento do template DOCX, inline de imagens e renderização final.

Camada 4 — único módulo que importa `docxtpl` e toca filesystem do template.

Origem: projeto SDD KPIReport (DS-02 / DS-05) — implementado em IM-03.
"""
from __future__ import annotations

import logging
import os
from io import BytesIO

from docx.opc.exceptions import PackageNotFoundError
from docx.shared import Mm  # docxtpl 0.20.2 não reexporta — usar python-docx (adaptação IM-03)
from docxtpl import DocxTemplate, InlineImage
from jinja2 import TemplateSyntaxError, UndefinedError

from src.utils import kpi_report_config as cfg

logger = logging.getLogger(__name__)


class TemplateLoadError(Exception):
    """Falha ao carregar ou renderizar template DOCX (SP-14).

    `self.path` = caminho tentado; `self.original` = exceção raiz.
    """

    def __init__(self, path: str, original: Exception) -> None:
        super().__init__(f"Falha no template DOCX: {path} ({type(original).__name__})")
        self.path = path
        self.original = original


def _resolve_template_path() -> str:
    """Resolve caminho final — override por env `KPI_REPORT_TEMPLATE_PATH` se arquivo existir;
    senão `TEMPLATE_PATH_DEFAULT` (SP-06 D6).
    """
    override = os.environ.get("KPI_REPORT_TEMPLATE_PATH", "").strip()
    if override and os.path.isfile(override):
        return override
    return cfg.TEMPLATE_PATH_DEFAULT


def _load_template() -> DocxTemplate:
    """Carrega `DocxTemplate` do caminho resolvido.

    Faz verificação eager (SP-14) — `DocxTemplate` é lazy (não valida no constructor),
    então checamos `os.path.isfile` antes. Converte qualquer falha para `TemplateLoadError`.
    """
    path = _resolve_template_path()

    if not os.path.isfile(path):
        raise TemplateLoadError(path, FileNotFoundError(f"Template não encontrado: {path}"))

    try:
        return DocxTemplate(path)
    except (PackageNotFoundError, KeyError) as exc:
        raise TemplateLoadError(path, exc) from exc


# Casas decimais por campo — espelha a formatação do PDF (kpi_report_v2_pdf._fmt = casas=1).
# O template DOCX imprime os valores crus via Jinja ({{ bloco1.kpis_planta.mttr }}), sem filtro,
# expondo floats como "31.362000000000002". Pré-formatamos aqui para strings pt-BR antes do render.
_DOCX_CASAS_KPI = {"mtbf": 1, "mttr": 1, "taxa_avaria": 1}
_DOCX_CASAS_LINHA = {
    "mtbf": 1, "mttr": 1, "taxa_avaria": 1,
    "avaria_min": 1, "atividade_h": 1, "uptime_h": 1,
    "num_ordens": 0, "num_paradas": 0,
}


def _fmt_campos(d: dict | None, casas_por_campo: dict[str, int]) -> None:
    """Formata in-place os campos numéricos de `d` conforme `casas_por_campo` (pt-BR)."""
    if not isinstance(d, dict):
        return
    for campo, casas in casas_por_campo.items():
        v = d.get(campo)
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            d[campo] = cfg.fmt_numero(v, casas=casas)


def _formatar_numeros_docx(dados: dict) -> None:
    """Pré-formata todos os números expostos ao template DOCX para strings pt-BR.

    Mutação in-place — `montar_docx` é o passo terminal do export DOCX, o dict não é
    reutilizado depois. Mesmo padrão da mutação de `sunburst_figures` (linha abaixo).
    """
    for bkey in ("bloco1", "bloco3"):
        bloco = dados.get(bkey)
        if not isinstance(bloco, dict):
            continue
        _fmt_campos(bloco.get("kpis_planta"), _DOCX_CASAS_KPI)
        _fmt_campos(bloco.get("metas"), _DOCX_CASAS_KPI)
        if isinstance(bloco.get("alert_range"), (int, float)) and not isinstance(bloco.get("alert_range"), bool):
            bloco["alert_range"] = cfg.fmt_numero(bloco["alert_range"], casas=1)

    bloco4 = dados.get("bloco4")
    if isinstance(bloco4, dict):
        for linha in (bloco4.get("linhas") or []):
            _fmt_campos(linha, _DOCX_CASAS_LINHA)
        _fmt_campos(bloco4.get("total"), _DOCX_CASAS_LINHA)

    for bkey in ("bloco2", "bloco5"):
        bloco = dados.get(bkey)
        if isinstance(bloco, dict):
            for parada in (bloco.get("paradas") or []):
                _fmt_campos(parada, {"duracao_min": 0})


def _inline_images_from_figures(
    template: DocxTemplate,
    figures_png: dict[str, bytes],
    width_mm: int = 80,
) -> dict[str, InlineImage]:
    """Converte dict `{kpi_key: bytes_png, ...}` em dict `{kpi_key: InlineImage, ...}`.

    Largura padrão 80 mm (DS-02 / DS-04); `docxtpl` preserva proporção automaticamente.
    """
    return {
        key: InlineImage(template, BytesIO(png_bytes), width=Mm(width_mm))
        for key, png_bytes in figures_png.items()
    }


def _mes_corrente_do_relatorio(dados: dict) -> tuple[int, str]:
    """Resolve `(ano, "YYYY-MM")` do mês corrente do relatório (SP-16 / DS-11).

    Preferência: início da janela mensal já calculada em `dados["periodo"]["inicio_mensal"]`
    — garante que o custo cobre o mesmo mês dos demais blocos. Fallback: `compute_monthly_window`
    sobre o agora no fuso configurado. Sem hardcode de data.
    """
    periodo = dados.get("periodo")
    inicio = periodo.get("inicio_mensal") if isinstance(periodo, dict) else None
    if inicio is None:
        inicio, _ = cfg.compute_monthly_window(cfg._now_in_report_timezone())
    return inicio.year, f"{inicio.year:04d}-{inicio.month:02d}"


def _inject_custo_mensal(template: DocxTemplate, dados: dict) -> None:
    """Injeta a figura de custo do mês corrente como InlineImage no contexto do template.

    DOCX-only (DS-11): roda no passo terminal do export, não em `coletar_dados_relatorio`.
    Defensivo — qualquer falha (sem dados, erro de import/Mongo de custos, figura `None`) deixa
    `dados['custo_mensal_img'] = None` e o template omite o bloco. Nunca derruba o DOCX.
    """
    try:
        ano, mes = _mes_corrente_do_relatorio(dados)
        from src.custos.figuras import figura_custo_mensal
        from src.utils.kpi_report_figures import renderizar_sunburst_png

        fig = figura_custo_mensal(ano, mes)
        if fig is None:
            dados["custo_mensal_img"] = None
            return
        # Render em alta resolução (mantém a proporção ~2.39:1 da aba) para nitidez ao
        # ocupar a largura cheia da página landscape.
        png = renderizar_sunburst_png(fig, kpi_label="custo_mensal",
                                      width=1600, height=668, scale=2)
        # 255 mm = largura útil cheia da página landscape (= largura das tabelas dos blocos,
        # 14480 twips); docxtpl preserva a proporção.
        dados["custo_mensal_img"] = InlineImage(template, BytesIO(png), width=Mm(255))
        dados["custo_mensal_titulo"] = f"Custo de Manutenção — {mes}"
    except Exception:
        logger.exception("Falha ao injetar gráfico de custo no DOCX; bloco omitido")
        dados["custo_mensal_img"] = None


_REC_LIMIT, _MAI_LIMIT = 10, 5


def _fmt_linha_lanc(l: dict) -> dict:
    """Formata uma linha de lançamento de custo para o template DOCX (SP-17 / DS-12).

    Data → `dd/mm/aaaa`; valor → `R$ x.xxx,xx` (pt-BR, 2 casas). Demais campos passam direto.
    """
    data = l.get("data")
    return {
        "data": data.strftime("%d/%m/%Y") if hasattr(data, "strftime") else "",
        "equipamento": l.get("equipamento", "") or "",
        "conta_nome": l.get("conta_nome", "") or "",
        "descritor": l.get("descritor", "") or "",
        "valor": "R$ " + cfg.fmt_numero(l.get("valor", 0) or 0, casas=2),
    }


def _inject_custo_tabelas(dados: dict) -> None:
    """Popula `custo_recentes` / `custo_maiores` (listas formatadas) no contexto do template.

    DOCX-only (DS-12) e defensivo — espelha `_inject_custo_mensal`. Qualquer falha (sem dados, erro
    de import/Mongo de custos) deixa as listas vazias e o template omite cada tabela. Nunca derruba
    o DOCX.
    """
    try:
        ano, mes = _mes_corrente_do_relatorio(dados)
        from src.custos.leitura import fetch_lancamentos_maiores, fetch_lancamentos_recentes

        rec = fetch_lancamentos_recentes(ano, mes, limit=_REC_LIMIT) or []
        mai = fetch_lancamentos_maiores(ano, mes, limit=_MAI_LIMIT) or []
        dados["custo_recentes"] = [_fmt_linha_lanc(l) for l in rec]
        dados["custo_maiores"] = [_fmt_linha_lanc(l) for l in mai]
    except Exception:
        logger.exception("Falha ao injetar tabelas de custo no DOCX; tabelas omitidas")
        dados["custo_recentes"] = []
        dados["custo_maiores"] = []


def montar_docx(dados: dict) -> bytes:
    """Carrega template, inline-a as imagens dos sunbursts, renderiza Jinja, retorna bytes do `.docx`.

    Captura `TemplateSyntaxError` / `UndefinedError` e re-lança como `TemplateLoadError` (SP-14).
    """
    template = _load_template()  # pode lançar TemplateLoadError

    # Envelopa PNGs dos sunbursts (Blocos 1 e 3) como InlineImage antes do render
    for bloco_key in ("bloco1", "bloco3"):
        bloco = dados.get(bloco_key)
        if not bloco:
            continue
        figs = bloco.get("sunburst_figures")
        if isinstance(figs, dict) and figs and isinstance(next(iter(figs.values())), (bytes, bytearray)):
            bloco["sunburst_figures"] = _inline_images_from_figures(template, figs)

    # Gráfico de Custo de Manutenção do mês corrente (SP-16 / DS-11) — DOCX-only, defensivo.
    _inject_custo_mensal(template, dados)

    # Tabelas de custo (últimos 10 + top 5 maiores) (SP-17 / DS-12) — DOCX-only, defensivo.
    _inject_custo_tabelas(dados)

    # Formata números crus (KPIs/metas/linhas/paradas) para strings pt-BR — o template
    # imprime via Jinja sem filtro, expondo floats como "31.362000000000002" (BR-07).
    _formatar_numeros_docx(dados)

    # BR-14 / SP-27: selo de parada pendente vai como texto na célula de causa (template intacto)
    _marcar_pendentes_docx(dados)

    try:
        template.render(dados, autoescape=True)
    except (TemplateSyntaxError, UndefinedError) as exc:
        raise TemplateLoadError(_resolve_template_path(), exc) from exc

    buffer = BytesIO()
    template.save(buffer)
    doc_bytes = _inserir_legenda_pendentes_docx(buffer.getvalue(), dados)
    return _inserir_mao_de_obra_docx(doc_bytes, dados)


# Texto puro no DOCX/PDF: o "⏳" depende da fonte de quem abre/imprime; o selo em texto basta.
_SELO_DOCX = "PENDENTE"
_BLOCOS_PARADAS = ("bloco2", "bloco5")  # ordem das tabelas "# | Equipamento | ..." no template


def _marcar_pendentes_docx(dados: dict) -> None:
    """Prefixa `PENDENTE · ` na causa das paradas ainda não consolidadas na NT0001 (BR-14)."""
    for bkey in _BLOCOS_PARADAS:
        for parada in ((dados.get(bkey) or {}).get("paradas") or []):
            if parada.get("pendente") and not str(parada.get("causa", "")).startswith(_SELO_DOCX):
                parada["causa"] = f"{_SELO_DOCX} · {parada.get('causa', '')}"


def _inserir_legenda_pendentes_docx(doc_bytes: bytes, dados: dict) -> bytes:
    """Legenda (se houver pendente) e aviso de checagem logo abaixo das tabelas dos Blocos 2 e 5.

    Pós-render via python-docx (template binário intacto). As tabelas de paradas são as que
    começam com "#" / "Equipamento", na ordem Bloco 2 → Bloco 5. Sem nada a inserir, devolve os
    bytes originais; qualquer falha também (nunca derruba o export).
    """
    textos = []
    for bkey in _BLOCOS_PARADAS:
        bloco = dados.get(bkey) or {}
        linhas = []
        if any(p.get("pendente") for p in (bloco.get("paradas") or [])):
            from src.utils.zpp_paradas_pendentes import LEGENDA
            linhas.append((LEGENDA.replace("⏳", "").strip(), False))
        if bloco.get("aviso_checagem"):
            linhas.append((bloco["aviso_checagem"], True))
        textos.append(linhas)
    if not any(textos):
        return doc_bytes
    try:
        from copy import deepcopy

        from docx import Document
        from docx.shared import Pt

        doc = Document(BytesIO(doc_bytes))
        tabelas = [t for t in doc.tables
                   if len(t.rows) and [c.text.strip() for c in t.rows[0].cells][:2] == ["#", "Equipamento"]]
        if len(tabelas) != len(_BLOCOS_PARADAS):
            logger.warning("DOCX: tabelas de paradas não identificadas (%d) — legenda omitida", len(tabelas))
            return doc_bytes
        for tabela, linhas in zip(tabelas, textos):
            ancora = tabela._tbl
            for texto, italico in linhas:
                par = doc.add_paragraph()
                run = par.add_run(texto)
                run.font.size = Pt(8)
                run.italic = italico
                elem = deepcopy(par._p)
                par._p.getparent().remove(par._p)
                ancora.addnext(elem)
                ancora = elem
        out = BytesIO()
        doc.save(out)
        return out.getvalue()
    except Exception:
        logger.exception("DOCX: falha ao inserir legenda de pendentes — export segue sem ela")
        return doc_bytes


def _inserir_mao_de_obra_docx(doc_bytes: bytes, dados: dict) -> bytes:
    """Insere o bloco "Mão de obra da semana" no DOCX já renderizado, ANTES do parágrafo
    "Anotações da Reunião" (SP-18/BR-11/12). Pós-render via python-docx — não edita o template
    binário. Defensivo: qualquer falha (HH indisponível, âncora não achada, kaleido) devolve o
    DOCX original intacto — nunca derruba o export (SP-19)."""
    try:
        from datetime import datetime as _dt
        from datetime import timedelta as _td

        from docx import Document
        from docx.shared import Mm, Pt

        from src.utils import kpi_report_config as _cfg
        from src.utils import kpi_report_maodeobra as _mdo
        from src.utils.hh_data import semana_mao_de_obra
        from src.utils.kpi_report_figures import renderizar_sunburst_png

        # dia avaliado = o mesmo do relatório (DatePicker, default ontem) — injetado pelo callback
        # de export em `dados["_mao_de_obra_ref"]`. Fallback: ontem.
        ref = dados.get("_mao_de_obra_ref")
        if ref is None:
            agora = _cfg._now_in_report_timezone() if hasattr(_cfg, "_now_in_report_timezone") else _dt.now()
            ref = agora - _td(days=1)
        sem = semana_mao_de_obra(ref)
        if not sem or sem.get("sem_dados") or not sem.get("dias"):
            return doc_bytes
        dias = sem["dias"]

        doc = Document(BytesIO(doc_bytes))
        anchor = next((p for p in doc.paragraphs if "Anotações da Reunião" in (p.text or "")), None)
        if anchor is None:
            logger.warning("DOCX mão de obra: âncora 'Anotações da Reunião' não encontrada; bloco omitido")
            return doc_bytes

        # heading + subtítulo com a semana de referência (deixa claro QUAL semana)
        h = doc.add_paragraph()
        run = h.add_run("Mão de obra da semana")
        run.bold = True
        run.font.size = Pt(14)
        sub = doc.add_paragraph()
        srun = sub.add_run(_mdo.rotulo_semana(sem))
        srun.italic = True
        srun.font.size = Pt(10)
        # imagem do gráfico (PNG kaleido — herda pin chromium IM-17)
        png = renderizar_sunburst_png(_mdo.fig_mao_de_obra(dias),
                                      kpi_label="mao_de_obra", width=1400, height=420, scale=2)
        doc.add_picture(BytesIO(png), width=Mm(255))
        pic_par = doc.paragraphs[-1]
        # tabela nativa
        linhas = _mdo.linhas_tabela(dias)
        tbl = doc.add_table(rows=1 + len(linhas), cols=len(_mdo.CABECALHOS))
        tbl.style = "Table Grid"
        for j, c in enumerate(_mdo.CABECALHOS):
            cel = tbl.rows[0].cells[j]
            cel.text = c
            for p in cel.paragraphs:
                for r in p.runs:
                    r.bold = True
        for i, l in enumerate(linhas, start=1):
            vals = (l["dia"], l["disponivel"], l["planejado"], l["pct"], l["executado"], l["aderencia"])
            for j, v in enumerate(vals):
                tbl.rows[i].cells[j].text = v

        # move heading → subtítulo → imagem → tabela para ANTES da âncora (mantém a ordem)
        anchor._p.addprevious(h._p)
        anchor._p.addprevious(sub._p)
        anchor._p.addprevious(pic_par._p)
        anchor._p.addprevious(tbl._tbl)

        out = BytesIO()
        doc.save(out)
        return out.getvalue()
    except Exception:
        logger.exception("Falha ao inserir bloco de mão de obra no DOCX; bloco omitido")
        return doc_bytes
